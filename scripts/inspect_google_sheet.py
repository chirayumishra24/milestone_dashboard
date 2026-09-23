import openpyxl

wb = openpyxl.load_workbook('google_sheet_downloaded.xlsx', data_only=True)
ws = wb['Sheet1']

print(f"Dimensions: {ws.max_row} rows x {ws.max_column} cols")

# Inspect Row 1 and Row 2 across ALL columns
row1 = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
row2 = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]

print("\n--- ALL COLUMNS (Row 1 & Row 2) ---")
current_exam = "General"
columns_meta = []

for idx in range(ws.max_column):
    c = idx + 1
    r1 = row1[idx]
    r2 = row2[idx]
    
    # If r1 has a title like Exam-1, Exam-2, Target, etc.
    if r1 is not None and str(r1).strip():
        current_exam = str(r1).strip()
    
    col_info = {
        'col_idx': c,
        'group_header': current_exam,
        'r1_raw': r1,
        'sub_header': str(r2).strip() if r2 is not None else None
    }
    columns_meta.append(col_info)
    print(f"Col {c:2d}: Group='{current_exam}' | SubHeader='{col_info['sub_header']}' | R1='{r1}'")

print("\n--- DATA ROWS SUMMARY ---")
student_count = 0
for r in range(3, ws.max_row + 1):
    sr_no = ws.cell(row=r, column=1).value
    enr_no = ws.cell(row=r, column=2).value
    name = ws.cell(row=r, column=3).value
    if name is not None and str(name).strip():
        student_count += 1
        non_empty = []
        for c in range(4, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None and str(val).strip():
                sub_h = ws.cell(row=2, column=c).value
                non_empty.append((c, sub_h, val))
        
        if student_count <= 5 or student_count == ws.max_row - 2:
            print(f"Student #{student_count}: Sr={sr_no} | Enr={enr_no} | Name={name}")
            print(f"   Non-empty count={len(non_empty)}: {non_empty[:8]}")

print(f"\nTotal Students detected: {student_count}")
