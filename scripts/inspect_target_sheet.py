import openpyxl

wb = openpyxl.load_workbook('CLASS IX TARGET SHEET.xlsx', data_only=True)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    headers = [ws.cell(row=1, column=c).value for c in range(1, 15)]
    count = 0
    names = []
    for r in range(2, ws.max_row + 1):
        n = ws.cell(row=r, column=2).value
        if n and str(n).strip():
            count += 1
            if count <= 3:
                names.append((r, n, [ws.cell(row=r, column=c).value for c in range(1, 11)]))
    print(f"\nSheet '{sheet}': Total students = {count}")
    print(f"Headers: {headers}")
    print("Sample rows:")
    for nr, n, vals in names:
        print(f"  Row {nr}: {vals}")
