import openpyxl

wb = openpyxl.load_workbook('CLASS IX TARGET SHEET.xlsx', data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    col9_nonempty = 0
    col10_nonempty = 0
    for r in range(2, ws.max_row + 1):
        v9 = ws.cell(row=r, column=9).value
        v10 = ws.cell(row=r, column=10).value
        if v9 is not None and str(v9).strip(): col9_nonempty += 1
        if v10 is not None and str(v10).strip(): col10_nonempty += 1
    print(f"Sheet {name}: col 9 ('OVERALL %') non-empty = {col9_nonempty}, col 10 ('TARGET GIVEN BY SCHOOL') non-empty = {col10_nonempty}")
