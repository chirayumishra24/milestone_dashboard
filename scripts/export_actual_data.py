"""
Export the Class IX dashboard data from the FMS workbook's "actual data" tab.

Reads ONLY the "actual data" tab and writes src/data/class9ActualData.ts (aggregate band
counts, no student records). Re-run whenever the sheet changes:

    python scripts/export_actual_data.py                      # the live Google Sheet (SHEET_URL)
    python scripts/export_actual_data.py "<google sheet link>"
    python scripts/export_actual_data.py "path/to/New Class IX FMS CCWS 26-27.xlsx"

A Google Sheet is downloaded through its xlsx export link, so it must be shared as
"anyone with the link can view".

Tab layout (one block per group, 20 rows each):
  - Block title in column A ("VIII" for the whole-class planner, "VIII - ENGLISH", ...).
  - Band rows: single bands ("90-94 %", "89-85 %", "33% and below") interleaved with
    cumulative "N and Above" rows. Only single bands are read; cumulative totals are
    recomputed by the dashboard because some cumulative cells in the sheet are wrong.
  - Columns read:  B = Class VIII half-yearly,  V = IX Pre-Mid Term actual,  AA = IX HY (Mid Term) actual.
    Target columns (U, Z, AC) are not exported.
"""
import datetime as dt
import json
import pathlib
import re
import sys
import tempfile
import urllib.parse
import urllib.request

import openpyxl
from openpyxl.utils import column_index_from_string

ROOT = pathlib.Path(__file__).resolve().parent.parent
# The live "New Class IX FMS CCWS 26-27" Google Sheet
SHEET_URL = 'https://docs.google.com/spreadsheets/d/1wh1uXUaWsZdhB_Yu1heT212WQbwqBc3drUDQwAt26S0/edit'
OUTPUT = ROOT / 'src' / 'data' / 'class9ActualData.ts'
TAB = 'actual data'

# Exam id -> sheet column holding that exam's student counts
EXAM_COLUMNS = {'viiiHalfYearly': 'B', 'preMid': 'V', 'midTerm': 'AA'}

# Lower bound of each single band -> band id used by the dashboard
RANGE_BANDS = {90: 'b90', 85: 'b85', 80: 'b80', 75: 'b75', 70: 'b70', 65: 'b65',
               60: 'b60', 55: 'b55', 50: 'b50', 45: 'b45', 40: 'b40', 33: 'b33'}
BAND_ORDER = ['b95', *RANGE_BANDS.values(), 'below33']
CUMULATIVE_MIN = {'95': 'b95', '90': 'b90', '80': 'b80', '70': 'b70', '60': 'b60', '50': 'b50', '40': 'b40'}

SUBJECTS = {
    'ENGLISH': ('english', 'English'),
    'HINDI': ('hindi', 'Hindi'),
    'SANSKRIT': ('sanskrit', 'Sanskrit'),
    'FRENCH': ('french', 'French'),
    'MATHS': ('maths', 'Mathematics'),
    'SCIENCE': ('science', 'Science'),
    'SOCIAL SCIENCE': ('socialScience', 'Social Science'),
    'COMPUTER': ('computer', 'Computer'),
}


def clean(label):
    # The sheet uses en dashes that come through as U+FFFD in some copies; normalise to "-"
    return re.sub(r'[‐-―�]', '-', str(label)).strip()


def classify(label):
    """Return ('band', id), ('cumulative', min_band_id) or None for a column-A label."""
    text = clean(label).lower()
    if re.match(r'^95\b', text) and 'above' in text:
        return 'band', 'b95'
    if 'below' in text and text.startswith('33'):
        return 'band', 'below33'
    m = re.match(r'^(\d+)\s*(?:%|marks)?\s*and above', text)
    if m:
        return 'cumulative', CUMULATIVE_MIN.get(m.group(1))
    m = re.match(r'^(\d+)\s*-\s*(\d+)', text)
    if m:
        low = min(int(m.group(1)), int(m.group(2)))
        if low in RANGE_BANDS:
            return 'band', RANGE_BANDS[low]
    return None


def block_identity(title):
    text = clean(title)
    if text.upper() == 'VIII':
        return 'wholeClass', 'Whole class'
    m = re.match(r'^VIII\s*-\s*(.+)$', text, re.I)
    if m and m.group(1).strip().upper() in SUBJECTS:
        return SUBJECTS[m.group(1).strip().upper()]
    raise ValueError(f'Unrecognised block title {title!r}')


def as_count(value, where):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    number = float(value)
    if not number.is_integer() or number < 0:
        raise ValueError(f'{where}: expected a whole student count, got {value!r}')
    return int(number)


def fetch_google_sheet(url):
    """Download a shared Google Sheet as xlsx. Returns (local path, sheet title)."""
    match = re.search(r'/spreadsheets/d/([\w-]+)', url)
    if not match:
        raise SystemExit(f'Not a Google Sheets link: {url}')
    export = f'https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=xlsx'
    try:
        with urllib.request.urlopen(export, timeout=120) as response:
            content_type = response.headers.get('Content-Type', '')
            disposition = response.headers.get('Content-Disposition', '')
            body = response.read()
    except OSError as error:
        raise SystemExit(f'Could not download the sheet ({error}). Check the link and your connection, '
                         'or pass a downloaded .xlsx file instead.')
    if 'spreadsheetml' not in content_type:
        raise SystemExit('Google did not return a spreadsheet. Share the sheet as "anyone with the link can view", '
                         'or pass a downloaded .xlsx file instead.')
    # e.g. filename*=UTF-8''New%20Class%20IX%20FMS%20CCWS%2026-27.xlsx
    name = re.search(r"filename\*=UTF-8''([^;]+)", disposition)
    title = urllib.parse.unquote(name.group(1)) if name else 'Google Sheet'
    tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    tmp.write(body)
    tmp.close()
    return pathlib.Path(tmp.name), title.removesuffix('.xlsx')


def main():
    source = sys.argv[1] if len(sys.argv) > 1 else SHEET_URL
    if source.startswith('http'):
        workbook_path, source_name = fetch_google_sheet(source)
        source_url = source
        saved_on = dt.date.today().isoformat()  # read live today
    else:
        workbook_path = pathlib.Path(source)
        source_name, source_url = workbook_path.name, None
        # The file's own timestamp: openpyxl's document "modified" property is not reliable
        saved_on = dt.date.fromtimestamp(workbook_path.stat().st_mtime).isoformat()
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    ws = wb[TAB]
    cols = {exam: column_index_from_string(letter) for exam, letter in EXAM_COLUMNS.items()}

    blocks, current, warnings = [], None, []
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label is None or not str(label).strip():
            continue
        kind = classify(label)
        if kind is None:
            gid, glabel = block_identity(label)
            current = {'id': gid, 'label': glabel, 'sheetTitle': clean(label), 'row': r,
                       'bands': {exam: {} for exam in EXAM_COLUMNS}, 'cumulative': {exam: {} for exam in EXAM_COLUMNS}}
            blocks.append(current)
            continue
        if current is None:
            raise ValueError(f'Row {r}: band row before any block title')
        what, band = kind
        for exam, c in cols.items():
            value = as_count(ws.cell(r, c).value, f'{TAB}!{EXAM_COLUMNS[exam]}{r}')
            if what == 'band':
                current['bands'][exam][band] = value
            elif band and value is not None:
                current['cumulative'][exam][band] = (value, f'{EXAM_COLUMNS[exam]}{r}')

    groups = []
    for block in blocks:
        results = {}
        for exam in EXAM_COLUMNS:
            bands = block['bands'][exam]
            missing = [b for b in BAND_ORDER if b not in bands]
            if missing:
                raise ValueError(f"{block['sheetTitle']}: band rows missing {missing}")
            if all(bands[b] is None for b in BAND_ORDER):
                continue  # this exam is not filled in for this group yet
            if any(bands[b] is None for b in BAND_ORDER):
                raise ValueError(f"{block['sheetTitle']} / {exam}: some band cells are empty")
            results[exam] = {b: bands[b] for b in BAND_ORDER}
            # Report (but do not use) cumulative cells that disagree with their band rows
            for min_band, (sheet_value, cell) in block['cumulative'][exam].items():
                computed = sum(bands[b] for b in BAND_ORDER[: BAND_ORDER.index(min_band) + 1])
                if computed != sheet_value:
                    warnings.append(f"{block['sheetTitle']} {cell}: sheet says {sheet_value}, band rows add up to {computed}")
        groups.append({'id': block['id'], 'label': block['label'], 'sheetTitle': block['sheetTitle'], 'results': results})

    whole = [g for g in groups if g['id'] == 'wholeClass']
    if len(whole) != 1:
        raise ValueError('Expected exactly one whole-class block')
    dataset = {
        'source': {
            'workbook': source_name,
            'tab': TAB,
            'sheetSavedOn': saved_on,
            **({'url': source_url} if source_url else {}),
        },
        'school': 'Cambridge Court World School',
        'academicYear': '2026-27',
        'wholeClass': whole[0],
        'subjects': [g for g in groups if g['id'] != 'wholeClass'],
    }

    body = json.dumps(dataset, indent=2, ensure_ascii=False)
    OUTPUT.write_text(
        '// Generated by scripts/export_actual_data.py from the "actual data" tab of\n'
        f'// "{source_name}". Do not edit by hand: update the sheet and re-run the script.\n'
        '// Aggregate student counts per score band only; no individual student records.\n'
        "import type { ActualDataset } from './milestoneBands';\n\n"
        f'export const CLASS_IX_ACTUAL_DATA: ActualDataset = {body};\n',
        encoding='utf-8',
    )

    print(f'Wrote {OUTPUT.relative_to(ROOT)} from "{source_name}" / {TAB}' + (f' ({source_url})' if source_url else ''))
    for g in groups:
        filled = ', '.join(f'{exam}={sum(counts.values())}' for exam, counts in g['results'].items()) or 'no data'
        print(f"  {g['label']:15s} {filled}")
    for w in warnings:
        print('  WARNING (sheet cumulative cell ignored):', w)


if __name__ == '__main__':
    main()
