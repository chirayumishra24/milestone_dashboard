import openpyxl
import json
import urllib.request
import sys

wb = openpyxl.load_workbook('CLASS IX TARGET SHEET.xlsx', data_only=True)
sheets = ['IX-AURA', 'IX-ZEN', 'IX-NEO']
all_students = []

for sheet_name in sheets:
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found!")
        continue
    
    sheet = wb[sheet_name]
    group = sheet_name.replace('IX-', '').strip()
    
    for r in range(2, sheet.max_row + 1):
        s_no = sheet.cell(r, 1).value
        name = sheet.cell(r, 2).value
        
        if not name or str(name).strip() == '':
            continue
            
        student_data = {
            'sNo': s_no if s_no is not None else (r - 1),
            'name': str(name).strip(),
            'group': group,
            'english': sheet.cell(r, 3).value,
            'maths': sheet.cell(r, 4).value,
            'sSt': sheet.cell(r, 5).value,
            'hsf': sheet.cell(r, 6).value,
            'science': sheet.cell(r, 7).value,
            'it': sheet.cell(r, 8).value,
            'overall': sheet.cell(r, 9).value,
            'target': sheet.cell(r, 10).value,
            'sourceRow': r,
            'sheetName': sheet_name
        }
        all_students.append(student_data)

print(f"Parsed {len(all_students)} students from Excel.")

# Send to API endpoint
url = 'http://localhost:3000/api/sync-students'
payload = json.dumps({
    'students': all_students,
    'syncSource': 'Local Excel Workbook Initial Ingestion'
}).encode('utf-8')

req = urllib.request.Request(
    url,
    data=payload,
    headers={
        'Content-Type': 'application/json',
        'x-sync-secret': 'ccis-alumni-sync-2026'
    },
    method='POST'
)

try:
    with urllib.request.urlopen(req) as response:
        res_body = response.read().decode('utf-8')
        print("API Response:", res_body)
except urllib.error.HTTPError as e:
    print(f"HTTP Error {e.code}: {e.read().decode('utf-8')}")
except Exception as e:
    print(f"Error connecting to {url}: {e}")
