import openpyxl

for fname in ['CLASS IX TARGET SHEET.xlsx', 'MASTER_CLASS_IX_MULTI_EXAM.xlsx']:
    print(f"\n==================== {fname} ====================")
    try:
        wb = openpyxl.load_workbook(fname, data_only=True)
        print(f"Sheet names: {wb.sheetnames}")
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"\n--- Sheet: {sheet_name} (Max row: {ws.max_row}, Max col: {ws.max_column}) ---")
            # Print top 5 rows
            for r in range(1, min(6, ws.max_row + 1)):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, min(25, ws.max_column + 1))]
                print(f"Row {r}: {row_vals}")
    except Exception as e:
        print(f"Error loading {fname}: {e}")
