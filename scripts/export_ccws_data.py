import json
import openpyxl
import os

roster_path = 'e:/skilizee/CCIS_Website/scripts/roster_5exam.json'
with open(roster_path, 'r', encoding='utf-8') as f:
    roster = json.load(f)

def compute_tier_stats(scores):
    valid = [float(s) for s in scores if s is not None]
    n = len(valid)
    if n == 0:
        return {
            'total': 0, 'average': 0.0, 'highest': 0.0, 'lowest': 0.0,
            'b95': 0, 'b95Pct': 0.0,
            'b90_94': 0, 'b90_94Pct': 0.0,
            'b90_plus': 0, 'b90_plusPct': 0.0,
            'b80_89': 0, 'b80_89Pct': 0.0,
            'b80_plus': 0, 'b80_plusPct': 0.0,
            'b70_79': 0, 'b70_79Pct': 0.0,
            'b70_plus': 0, 'b70_plusPct': 0.0,
            'b60_69': 0, 'b60_69Pct': 0.0,
            'b60_plus': 0, 'b60_plusPct': 0.0,
            'b50_59': 0, 'b50_59Pct': 0.0,
            'below50': 0, 'below50Pct': 0.0,
            'below60': 0, 'below60Pct': 0.0,
        }
    b95 = sum(1 for s in valid if s >= 95.0)
    b90_94 = sum(1 for s in valid if 90.0 <= s < 95.0)
    b90_plus = b95 + b90_94
    b80_89 = sum(1 for s in valid if 80.0 <= s < 90.0)
    b80_plus = b90_plus + b80_89
    b70_79 = sum(1 for s in valid if 70.0 <= s < 80.0)
    b70_plus = b80_plus + b70_79
    b60_69 = sum(1 for s in valid if 60.0 <= s < 70.0)
    b60_plus = b70_plus + b60_69
    b50_59 = sum(1 for s in valid if 50.0 <= s < 60.0)
    below50 = sum(1 for s in valid if s < 50.0)
    below60 = sum(1 for s in valid if s < 60.0)
    avg = sum(valid) / n
    return {
        'total': n,
        'average': round(avg, 1),
        'highest': round(max(valid), 1),
        'lowest': round(min(valid), 1),
        'b95': b95,
        'b95Pct': round((b95 / n) * 100.0, 1),
        'b90_94': b90_94,
        'b90_94Pct': round((b90_94 / n) * 100.0, 1),
        'b90_plus': b90_plus,
        'b90_plusPct': round((b90_plus / n) * 100.0, 1),
        'b80_89': b80_89,
        'b80_89Pct': round((b80_89 / n) * 100.0, 1),
        'b80_plus': b80_plus,
        'b80_plusPct': round((b80_plus / n) * 100.0, 1),
        'b70_79': b70_79,
        'b70_79Pct': round((b70_79 / n) * 100.0, 1),
        'b70_plus': b70_plus,
        'b70_plusPct': round((b70_plus / n) * 100.0, 1),
        'b60_69': b60_69,
        'b60_69Pct': round((b60_69 / n) * 100.0, 1),
        'b60_plus': b60_plus,
        'b60_plusPct': round((b60_plus / n) * 100.0, 1),
        'b50_59': b50_59,
        'b50_59Pct': round((b50_59 / n) * 100.0, 1),
        'below50': below50,
        'below50Pct': round((below50 / n) * 100.0, 1),
        'below60': below60,
        'below60Pct': round((below60 / n) * 100.0, 1),
    }

# Overall
e1_ov = [s.get('exams',{}).get('exam-1',{}).get('overall',{}).get('value') for s in roster]
e2_ov = [s.get('exams',{}).get('exam-2',{}).get('overall',{}).get('value') for s in roster]
tgt_ov = [s.get('schoolTarget',{}).get('overall',{}).get('value') for s in roster if s.get('schoolTarget')]

overall = {
    'exam1': compute_tier_stats(e1_ov),
    'exam2': compute_tier_stats(e2_ov),
    'target': compute_tier_stats(tgt_ov),
}

# Sections
sections = ['AURA', 'ZEN', 'NEO']
section_data = {}
for sec in sections:
    sec_roster = [s for s in roster if s.get('group') == sec]
    e1_sec = [s.get('exams',{}).get('exam-1',{}).get('overall',{}).get('value') for s in sec_roster]
    e2_sec = [s.get('exams',{}).get('exam-2',{}).get('overall',{}).get('value') for s in sec_roster]
    tgt_sec = [s.get('schoolTarget',{}).get('overall',{}).get('value') for s in sec_roster if s.get('schoolTarget')]
    section_data[sec] = {
        'sectionName': f'IX-{sec}',
        'totalEnrolled': len(sec_roster),
        'exam1': compute_tier_stats(e1_sec),
        'exam2': compute_tier_stats(e2_sec),
        'target': compute_tier_stats(tgt_sec)
    }

# Subjects
subjects_config = [
    {'id': 'english', 'label': 'English Language & Literature', 'shortLabel': 'English', 'code': 'ENG 184'},
    {'id': 'maths', 'label': 'Mathematics Standard', 'shortLabel': 'Mathematics', 'code': 'MATH 041'},
    {'id': 'science', 'label': 'Science (Theory & Lab)', 'shortLabel': 'Science', 'code': 'SCI 086'},
    {'id': 'socialScience', 'label': 'Social Science', 'shortLabel': 'Social Science', 'code': 'SST 087'},
    {'id': 'it', 'label': 'Information Technology', 'shortLabel': 'Info Tech (IT)', 'code': 'IT 402'},
    {'id': 'secondLanguage', 'label': 'Second Language (Combined H/S/F)', 'shortLabel': '2nd Language', 'code': 'LANG-2'},
]

subjects_data = {}
for sub in subjects_config:
    sub_id = sub['id']
    e1_sub = []
    e2_sub = []
    tgt_sub = []
    for s in roster:
        v1 = s.get('exams',{}).get('exam-1',{}).get('subjects',{}).get(sub_id,{}).get('value')
        if v1 is not None:
            e1_sub.append((v1 / 20.0) * 100.0)
        v2 = s.get('exams',{}).get('exam-2',{}).get('subjects',{}).get(sub_id,{}).get('value')
        if v2 is not None:
            e2_sub.append((v2 / 20.0) * 100.0)
        t_obj = s.get('schoolTarget') or {}
        st_sub = (t_obj.get('subjects') or {}).get(sub_id, {}).get('value')
        if st_sub is not None:
            tgt_sub.append(float(st_sub))
    subjects_data[sub_id] = {
        'id': sub['id'],
        'label': sub['label'],
        'shortLabel': sub['shortLabel'],
        'code': sub['code'],
        'exam1': compute_tier_stats(e1_sub),
        'exam2': compute_tier_stats(e2_sub),
        'target': compute_tier_stats(tgt_sub) if tgt_sub else None
    }

# Second language breakdown for Exam 2
languages = [
    {'id': 'hindi', 'label': 'Hindi Course-A', 'code': 'HIN 002', 'name': 'Hindi'},
    {'id': 'sanskrit', 'label': 'Sanskrit Communicative', 'code': 'SAN 122', 'name': 'Sanskrit'},
    {'id': 'french', 'label': 'French Language', 'code': 'FRN 018', 'name': 'French'},
]
language_data = {}
for l in languages:
    scores = []
    for s in roster:
        if s.get('secondLanguage') == l['name']:
            v = s.get('exams',{}).get('exam-2',{}).get('subjects',{}).get('secondLanguage',{}).get('value')
            if v is not None:
                scores.append((v / 20.0) * 100.0)
    language_data[l['id']] = {
        'id': l['id'],
        'label': l['label'],
        'code': l['code'],
        'name': l['name'],
        'exam2': compute_tier_stats(scores)
    }

# FMS Benchmark
wb = openpyxl.load_workbook('e:/skilizee/milestone/scripts/source-data/New Class IX FMS CCWS 26-27.xlsx', data_only=True)
step2 = wb['STEP 2']
fms_rows = []
for r in range(4, 24):
    b = str(step2.cell(r, 1).value or '').strip()
    fms_rows.append({
        'bracket': b,
        'class8HY': step2.cell(r, 2).value or 0,
        'preMidTarget': step2.cell(r, 10).value or 0,
        'preMidActual': step2.cell(r, 11).value or 0,
        'midTermTarget': step2.cell(r, 15).value or 0,
        'finalTarget': step2.cell(r, 18).value or 0,
    })

header_ts = """/**
 * Class IX Academic Performance & Score Tier Analytics
 * Institution: Cambridge Court World School (CCWS)
 * Academic Session: 2026–27
 * Note: Aggregate statistics only — strictly zero individual scholar personal records.
 */

export interface ScoreTierDistribution {
  total: number;
  average: number;
  highest: number;
  lowest: number;
  b95: number;
  b95Pct: number;
  b90_94: number;
  b90_94Pct: number;
  b90_plus: number;
  b90_plusPct: number;
  b80_89: number;
  b80_89Pct: number;
  b80_plus: number;
  b80_plusPct: number;
  b70_79: number;
  b70_79Pct: number;
  b70_plus: number;
  b70_plusPct: number;
  b60_69: number;
  b60_69Pct: number;
  b60_plus: number;
  b60_plusPct: number;
  b50_59: number;
  b50_59Pct: number;
  below50: number;
  below50Pct: number;
  below60: number;
  below60Pct: number;
}

export interface SubjectPerformanceItem {
  id: string;
  label: string;
  shortLabel: string;
  code: string;
  exam1: ScoreTierDistribution;
  exam2: ScoreTierDistribution;
  target: ScoreTierDistribution | null;
}

export interface SectionPerformanceItem {
  sectionName: string;
  totalEnrolled: number;
  exam1: ScoreTierDistribution;
  exam2: ScoreTierDistribution;
  target: ScoreTierDistribution;
}

export interface SecondLanguageItem {
  id: string;
  label: string;
  code: string;
  name: string;
  exam2: ScoreTierDistribution;
}

export interface FmsBenchmarkRow {
  bracket: string;
  class8HY: number;
  preMidTarget: number;
  preMidActual: number;
  midTermTarget: number;
  finalTarget: number;
}
"""

school_dict = {
    'name': 'Cambridge Court World School',
    'shortName': 'CCWS',
    'affiliation': 'CBSE Affiliated No. 1730867',
    'cohort': 'Class IX Secondary',
    'session': '2026–27',
    'totalScholarsEnrolled': len(roster),
    'sections': ['IX-AURA', 'IX-ZEN', 'IX-NEO'],
    'lastUpdated': 'September 2026',
}

full_ts = header_ts + "\nexport const CLASS_IX_PERFORMANCE_METRICS = " + json.dumps({
    'school': school_dict,
    'overall': overall,
    'sections': section_data,
    'subjects': subjects_data,
    'electives': language_data,
    'fmsBenchmark': fms_rows
}, indent=2) + ";\n"

out_path = 'e:/skilizee/CCWS-Website/src/data/class9PerformanceData.ts'
os.makedirs(os.path.dirname(out_path), exist_ok=True)
with open(out_path, 'w', encoding='utf-8') as f:
    f.write(full_ts)

print(f"SUCCESS: Written {len(full_ts)} bytes to {out_path}")
