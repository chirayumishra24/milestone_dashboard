import openpyxl

wb = openpyxl.load_workbook("New Class IX FMS CCWS 26-27.xlsx", data_only=True)
ws_fms = wb['FMS']

print("FMS Steps:")
current_step = None
for r in range(1, 60):
    val_a = ws_fms.cell(row=r, column=1).value
    val_b = ws_fms.cell(row=r, column=2).value
    val_c = ws_fms.cell(row=r, column=3).value
    val_d = ws_fms.cell(row=r, column=4).value
    
    if val_a is not None and str(val_a).strip():
        current_step = str(val_a).strip()
        print(f"\n[{current_step}]")
    
    if val_b or val_c or val_d:
        print(f"  {str(val_b).strip() if val_b else ''}: {str(val_c).strip() if val_c else ''} | {str(val_d).strip() if val_d else ''}")
