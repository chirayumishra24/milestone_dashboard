import openpyxl

wb = openpyxl.load_workbook('MASTER_CLASS_IX_MULTI_EXAM.xlsx', data_only=True)
ws = wb['Sheet1']

exams = ['Exam-1', 'Exam-2', 'Exam-3', 'Exam-4', 'Exam-5', 'Target']
cols_by_group = {}
r1 = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
r2 = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]

cur_g = None
for c in range(1, ws.max_column + 1):
    val1 = r1[c-1]
    val2 = r2[c-1]
    if val1 and str(val1).strip():
        cur_g = str(val1).strip()
    if cur_g not in cols_by_group:
        cols_by_group[cur_g] = []
    cols_by_group[cur_g].append((c, str(val2).strip() if val2 else ''))

print("Exam Groups and Columns:")
for g, cols in cols_by_group.items():
    print(f"\n{g}:")
    for c, sub in cols:
        print(f"  Col {c}: {sub}")

# Check non-empty rows count for each exam
print("\n--- Data Presence Per Exam ---")
for g in ['Exam-1', 'Exam-2', 'Exam-3', 'Exam-4', 'Exam-5', 'Target']:
    cols = [c for c, sub in cols_by_group.get(g, [])]
    non_empty_students = 0
    for r in range(3, ws.max_row + 1):
        has_val = any(ws.cell(row=r, column=c).value is not None and str(ws.cell(row=r, column=c).value).strip() != '' for c in cols)
        if has_val:
            non_empty_students += 1
    print(f"{g}: {non_empty_students} students have data out of 97")

# Check section breakdown
sections = {}
for r in range(3, ws.max_row + 1):
    enr = ws.cell(row=r, column=2).value
    name = ws.cell(row=r, column=3).value
    if name and str(name).strip():
        sec = "Unknown"
        if enr and "AURA" in str(enr): sec = "IX-AURA"
        elif enr and "ZEN" in str(enr): sec = "IX-ZEN"
        elif enr and "NEO" in str(enr): sec = "IX-NEO"
        sections[sec] = sections.get(sec, 0) + 1

print("\nSection breakdown in MASTER_CLASS_IX_MULTI_EXAM:", sections)
