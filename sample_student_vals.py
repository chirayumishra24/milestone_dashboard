import openpyxl

wb = openpyxl.load_workbook('MASTER_CLASS_IX_MULTI_EXAM.xlsx', data_only=True)
ws = wb['Sheet1']

print("Sample Student 1 (Row 3):")
for c in range(1, ws.max_column + 1):
    g = ws.cell(row=1, column=c).value
    s = ws.cell(row=2, column=c).value
    v = ws.cell(row=3, column=c).value
    if v is not None and str(v).strip():
        print(f"  Col {c:2d} | Group: {str(g):12s} | Sub: {str(s):15s} | Val: {v}")

print("\nSample Student 2 (Row 4):")
for c in range(1, ws.max_column + 1):
    g = ws.cell(row=1, column=c).value
    s = ws.cell(row=2, column=c).value
    v = ws.cell(row=4, column=c).value
    if v is not None and str(v).strip():
        print(f"  Col {c:2d} | Group: {str(g):12s} | Sub: {str(s):15s} | Val: {v}")
