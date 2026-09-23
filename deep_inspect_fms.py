import openpyxl

wb = openpyxl.load_workbook("New Class IX FMS CCWS 26-27.xlsx", data_only=True)

# 1. Inspect FMS Sheet completely
ws_fms = wb['FMS']
print("="*30 + " COMPLETE FMS SHEET " + "="*30)
for r in range(1, ws_fms.max_row + 1):
    vals = [ws_fms.cell(row=r, column=c).value for c in range(1, 10)]
    if any(v is not None and str(v).strip() for v in vals):
        print(f"Row {r:2d}: {vals[:5]}")

# 2. Inspect STEP 1 TARGET SHEET completely
ws_step1 = wb['STEP 1 TARGET SHEET']
print("\n" + "="*30 + " STEP 1 TARGET SHEET " + "="*30)
headers = [ws_step1.cell(row=2, column=c).value for c in range(1, ws_step1.max_column + 1)]
print("Headers (Row 2):", headers)

students = []
for r in range(3, ws_step1.max_row + 1):
    sr = ws_step1.cell(row=r, column=1).value
    enr = ws_step1.cell(row=r, column=2).value
    name = ws_step1.cell(row=r, column=3).value
    gender = ws_step1.cell(row=r, column=4).value
    if name is not None and str(name).strip():
        # Get all subject columns
        row_data = [ws_step1.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        students.append(row_data)

print(f"Total students found in STEP 1 TARGET SHEET: {len(students)}")
print("First 3 students:")
for s in students[:3]:
    print(s)
print("Last 3 students:")
for s in students[-3:]:
    print(s)

# 3. Inspect STEP 2
ws_step2 = wb['STEP 2']
print("\n" + "="*30 + " STEP 2 HEADERS & SAMPLE ROWS " + "="*30)
for r in range(1, 25):
    vals = [ws_step2.cell(row=r, column=c).value for c in range(1, ws_step2.max_column + 1)]
    if any(v is not None and str(v).strip() for v in vals):
        print(f"Row {r:2d}: {vals[:15]}")

# 4. Inspect actual data
ws_actual = wb['actual data']
print("\n" + "="*30 + " ACTUAL DATA SAMPLE ROWS " + "="*30)
for r in range(1, 25):
    vals = [ws_actual.cell(row=r, column=c).value for c in range(1, ws_actual.max_column + 1)]
    if any(v is not None and str(v).strip() for v in vals):
        print(f"Row {r:2d}: {vals[:15]}")
