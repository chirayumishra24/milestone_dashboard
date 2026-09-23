import openpyxl

wb = openpyxl.load_workbook('MASTER_CLASS_IX_MULTI_EXAM.xlsx', data_only=True)
ws = wb['Sheet1']

print(f"MASTER_CLASS_IX_MULTI_EXAM: {ws.max_row} rows, {ws.max_column} cols")
r1 = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
r2 = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]

cur_group = "General"
for c in range(1, ws.max_column + 1):
    val1 = r1[c-1]
    val2 = r2[c-1]
    if val1 is not None and str(val1).strip():
        cur_group = str(val1).strip()
    print(f"Col {c:2d} | Group: {cur_group:15s} | R2: {str(val2):20s} | R1_raw: {str(val1)}")

print("\n--- SAMPLE ROWS ---")
for r in range(3, min(10, ws.max_row + 1)):
    sr = ws.cell(row=r, column=1).value
    enr = ws.cell(row=r, column=2).value
    name = ws.cell(row=r, column=3).value
    print(f"Row {r}: Sr={sr}, Enr={enr}, Name={name}")

# Total non-empty student rows
count = 0
for r in range(3, ws.max_row + 1):
    name = ws.cell(row=r, column=3).value
    if name and str(name).strip():
        count += 1
print(f"\nTotal valid student rows: {count}")
