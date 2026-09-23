import openpyxl
import json

wb = openpyxl.load_workbook('CLASS IX TARGET SHEET.xlsx', data_only=True)
sheets = ['IX-AURA', 'IX-ZEN', 'IX-NEO']
all_students = []

for s in sheets:
    sheet = wb[s]
    group = s.replace('IX-', '').strip()
    for r in range(2, sheet.max_row + 1):
        s_no = sheet.cell(r, 1).value
        name = sheet.cell(r, 2).value
        if not name or str(name).strip() == '':
            continue
        all_students.append({
            'sNo': s_no if s_no is not None else (r - 1),
            'name': str(name).strip(),
            'group': group,
            'english': sheet.cell(r, 3).value,
            'maths': sheet.cell(r, 4).value,
            'sSt': sheet.cell(r, 5).value,
            'hsf': sheet.cell(r, 6).value,
            'science': sheet.cell(r, 7).value,
            'it': sheet.cell(r, 8).value,
            'overall': sheet.cell(r, 9).value,
            'target': sheet.cell(r, 10).value,
            'sourceRow': r,
            'sheetName': s
        })

with open('scripts/students.json', 'w', encoding='utf-8') as f:
    json.dump(all_students, f, indent=2)

print(f"Successfully extracted {len(all_students)} student rows to scripts/students.json")
